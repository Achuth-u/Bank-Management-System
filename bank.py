from openpyxl import Workbook, load_workbook
import os
import hashlib
import random
from datetime import datetime



# CREATE EXCEL FILES


if not os.path.exists("bank.xlsx"):
    wb = Workbook()

    ws = wb.active
    ws.title = "Accounts"

    ws.append([
        "Account No",
        "Name",
        "Age",
        "Phone",
        "Password Hash",
        "Balance"
    ])

    wb.create_sheet("Transactions")   
    ts = wb["Transactions"]

    ts.append([
        "Transaction ID",
        "Account No",
        "Type",
        "Amount",
        "Balance",
        "Date"
    ])

    wb.save("bank.xlsx")


# PASSWORD HASH

def hash_password(password):
    return hashlib.sha256(password.encode()).hexdigest()

# GENERATE ACCOUNT NUMBER

def generate_account_number():

    wb = load_workbook("bank.xlsx")
    ws = wb["Accounts"]

    while True:

        account_no = random.randint(100000, 999999)

        found = False

        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == account_no:
                found = True
                break

        if not found:
            return account_no


# CREATE ACCOUNT

def create_account():

    wb = load_workbook("bank.xlsx")
    ws = wb["Accounts"]

    print("\n CREATE ACCOUNT ")

    name = input("Enter name: ")
    age = int(input("Enter age: "))
    phone = input("Enter phone number: ")

    password = input("Create password: ")

    account_no = generate_account_number()

    password_hash = hash_password(password)

    balance = 0

    ws.append([
        account_no,
        name,
        age,
        phone,
        password_hash,
        balance
    ])

    wb.save("bank.xlsx")

    print("\nAccount created successfully!")
    print("Your Account Number:", account_no)



# LOGIN

def login():

    wb = load_workbook("bank.xlsx")
    ws = wb["Accounts"]

    print("\n LOGIN ")

    account_no = int(input("Enter account number: "))
    password = input("Enter password: ")

    password_hash = hash_password(password)

    for row in ws.iter_rows(min_row=2):

        if row[0].value == account_no:

            if row[4].value == password_hash:

                print("\nLogin successful!")

                return account_no

            else:
                print("Incorrect password.")
                return None

    print("Account not found.")
    return None



# CHECK BALANCE

def check_balance(account_no):

    wb = load_workbook("bank.xlsx")
    ws = wb["Accounts"]

    for row in ws.iter_rows(min_row=2):

        if row[0].value == account_no:

            print("\nCurrent Balance:", row[5].value)
            return



# DEPOSIT

def deposit(account_no):

    wb = load_workbook("bank.xlsx")

    ws = wb["Accounts"]
    ts = wb["Transactions"]

    amount = float(input("Enter deposit amount: "))

    if amount <= 0:
        print("Invalid amount.")
        return

    for row in ws.iter_rows(min_row=2):

        if row[0].value == account_no:

            old_balance = row[5].value or 0
            new_balance = old_balance + amount

            row[5].value = new_balance

            transaction_id = "TX" + str(random.randint(100000, 999999))

            ts.append([
                transaction_id,
                account_no,
                "Deposit",
                amount,
                new_balance,
                datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            ])

            wb.save("bank.xlsx")

            print("Deposit successful.")
            print("New Balance:", new_balance)

            return

# WITHDRAW

def withdraw(account_no):

    wb = load_workbook("bank.xlsx")

    ws = wb["Accounts"]
    ts = wb["Transactions"]

    amount = float(input("Enter withdrawal amount: "))

    for row in ws.iter_rows(min_row=2):

        if row[0].value == account_no:

            old_balance = row[5].value or 0

            if amount <= 0:
                print("Invalid amount.")
                return

            if amount > old_balance:
                print("Insufficient balance.")
                return

            new_balance = old_balance - amount

            row[5].value = new_balance

            transaction_id = "TX" + str(random.randint(100000, 999999))

            ts.append([
                transaction_id,
                account_no,
                "Withdrawal",
                amount,
                new_balance,
                datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            ])

            wb.save("bank.xlsx")

            print("Withdrawal successful.")
            print("Remaining Balance:", new_balance)

            return


# TRANSACTION HISTORY

def transaction_history(account_no):

    wb = load_workbook("bank.xlsx")
    ts = wb["Transactions"]

    print("\n    TRANSACTION HISTORY    ")

    found = False

    for row in ts.iter_rows(min_row=2, values_only=True):

        if row[1] == account_no:

            found = True

            print(
                row[0],
                "|",
                row[2],
                "|",
                row[3],
                "| Balance:",
                row[4],
                "|",
                row[5]
            )

    if not found:
        print("No transactions found.")


# ACCOUNT MENU

def account_menu(account_no):

    while True:

        print("\n===== BANK MENU =====")

        print("1. Check Balance")
        print("2. Deposit")
        print("3. Withdraw")
        print("4. Transaction History")
        print("5. Logout")

        choice = input("Enter choice: ")

        if choice == "1":
            check_balance(account_no)

        elif choice == "2":
            deposit(account_no)

        elif choice == "3":
            withdraw(account_no)

        elif choice == "4":
            transaction_history(account_no)

        elif choice == "5":
            print("Logged out.")
            break

        else:
            print("Invalid choice.")


# MAIN PROGRAM

while True:

    print("     BANK MANAGEMENT SYSTEM")

    print("1. Create Account")
    print("2. Login")
    print("3. Exit")

    choice = input("Enter choice: ")

    if choice == "1":

        create_account()

    elif choice == "2":

        account_no = login()

        if account_no:
            account_menu(account_no)

    elif choice == "3":

        print("Thank you!")
        break

    else:

        print("Invalid choice.")